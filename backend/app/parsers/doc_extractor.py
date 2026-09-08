import io
import base64
from typing import Dict, Any, Optional
from pypdf import PdfReader
import docx
from PIL import Image

class ExtractedDocument:
    def __init__(
        self,
        filename: str,
        content_type: str,
        raw_text: str = "",
        image_base64: Optional[str] = None,
        detected_type: str = "General Document",
        metadata: Optional[Dict[str, Any]] = None
    ):
        self.filename = filename
        self.content_type = content_type
        self.raw_text = raw_text
        self.image_base64 = image_base64
        self.detected_type = detected_type
        self.metadata = metadata or {}

    def to_dict(self) -> Dict[str, Any]:
        return {
            "filename": self.filename,
            "content_type": self.content_type,
            "detected_type": self.detected_type,
            "char_count": len(self.raw_text),
            "snippet": self.raw_text[:300] if self.raw_text else "[Binary/Image content]",
            "metadata": self.metadata
        }

def classify_document(text: str, filename: str) -> str:
    fn_lower = filename.lower()
    text_lower = text.lower()
    combined = f"{fn_lower} {text_lower}"

    # 1. Titer / Serology Lab Report (Prioritize before Rabies, because titer reports always contain 'rabies')
    if any(k in fn_lower for k in ["titer", "favn", "rnatt", "serolog"]) or \
       any(k in text_lower for k in ["titer", "favn", "rnatt", "serology", "antibody titration", "fluorescent antibody"]):
        return "Rabies Titer (FAVN/RNATT) Lab Report"

    # 2. EU Annex IV Health Certificate
    if any(k in fn_lower for k in ["annex_iv", "annex-iv", "annex iv", "annex4", "annex_4"]) or \
       any(k in text_lower for k in ["annex iv", "annex 4", "model animal health certificate", "non-commercial movement into a member state"]):
        return "EU Annex IV Health Certificate"

    # 3. Owner Non-Commercial Declaration
    if any(k in fn_lower for k in ["declaration", "owner_dec", "owner-dec", "non_commercial", "non-commercial"]) or \
       any(k in text_lower for k in ["owner declaration", "non-commercial declaration", "declaration of owner", "declaration of non-commercial", "5-day window", "5 days of the movement"]):
        return "Non-Commercial Owner Declaration"

    # 4. Official Pet Passport
    if any(k in fn_lower for k in ["passport", "passeport", "pet_pass"]) or \
       any(k in text_lower for k in ["pet passport", "passeport pour animaux", "official pet passport", "veterinary passport"]):
        return "Official Pet Passport"

    # 5. Official Veterinary Health Certificate (General/Non-Annex)
    if any(k in fn_lower for k in ["health_cert", "health-cert", "healthcert", "vet_cert", "ahc"]) or \
       any(k in text_lower for k in ["veterinary health certificate", "animal health certificate", "official health certificate", "certificate of veterinary inspection", "fit to travel", "clinical examination"]):
        return "Official Veterinary Health Certificate"

    # 6. Government Export / Endorsement Permit
    if any(k in fn_lower for k in ["export", "permit", "daff", "endorsement", "import_permit"]) or \
       any(k in text_lower for k in ["export permit", "notice of intention", "daff", "usda endorsed", "aphis form", "import permit"]):
        return "Government Export / Endorsement Permit"

    # 7. Internal Parasite / Tapeworm Treatment Record
    if any(k in fn_lower for k in ["tapeworm", "echinococcus", "worm", "deworm", "parasite"]) or \
       any(k in text_lower for k in ["echinococcus", "praziquantel", "tapeworm", "parasite treatment"]):
        return "Internal Parasite / Tapeworm Treatment Record"

    # 8. Microchip Registration Record
    if any(k in fn_lower for k in ["microchip", "chip", "transponder", "iso11784"]) or \
       any(k in text_lower for k in ["transponder implantation", "microchip registration", "petlog", "avid", "homeagain", "identichip", "iso 11784", "iso 11785"]):
        return "Microchip Registration Record"

    # 9. Rabies / Vaccination Certificate
    if any(k in fn_lower for k in ["rabies", "vaccin", "vax", "immunis"]) or \
       any(k in text_lower for k in ["rabies", "vaccination", "immunisation", "rabisin", "defensor", "nobivac rabies", "dhpp"]):
        return "Rabies / Vaccination Certificate"

    # 10. Fallbacks based on weaker keyword occurrences
    if any(k in combined for k in ["microchip", "transponder"]):
        return "Microchip Registration Record"
    if any(k in combined for k in ["health certificate", "certificate"]):
        return "Official Veterinary Health Certificate"

    return "General Pet Record / Photo"

def extract_from_bytes(file_bytes: bytes, filename: str, content_type: str) -> ExtractedDocument:
    ext = filename.split(".")[-1].lower() if "." in filename else ""
    raw_text = ""
    image_b64 = None
    meta: Dict[str, Any] = {"file_size_bytes": len(file_bytes)}

    # 1. PDF
    if ext == "pdf" or "pdf" in content_type:
        try:
            reader = PdfReader(io.BytesIO(file_bytes))
            meta["page_count"] = len(reader.pages)
            pages_text = []
            for i, page in enumerate(reader.pages):
                txt = page.extract_text() or ""
                pages_text.append(txt)
            raw_text = "\n\n".join(pages_text)
        except Exception as e:
            raw_text = f"[PDF read error: {str(e)}]"

    # 2. DOCX / Word
    elif ext in ["docx", "doc"] or "word" in content_type:
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            para_text = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    para_text.append(" | ".join(cell.text.strip() for cell in row.cells))
            raw_text = "\n".join(para_text)
        except Exception as e:
            raw_text = f"[Word document read error: {str(e)}]"

    # 3. Plain Text
    elif ext in ["txt", "csv", "json"] or "text" in content_type:
        try:
            raw_text = file_bytes.decode("utf-8", errors="ignore")
        except Exception:
            raw_text = ""

    # 4. Images (JPG, PNG, WEBP)
    elif ext in ["jpg", "jpeg", "png", "webp"] or "image" in content_type:
        try:
            img = Image.open(io.BytesIO(file_bytes))
            meta["image_width"] = img.width
            meta["image_height"] = img.height
            meta["image_format"] = img.format
            image_b64 = base64.b64encode(file_bytes).decode("utf-8")
        except Exception as e:
            meta["image_error"] = str(e)

    detected_type = classify_document(raw_text, filename)
    return ExtractedDocument(
        filename=filename,
        content_type=content_type,
        raw_text=raw_text,
        image_base64=image_b64,
        detected_type=detected_type,
        metadata=meta
    )
